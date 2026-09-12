"""Build the Persian Word documentation for nevox news!."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


def set_run_font(run, size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn('w:ascii'), 'Calibri')
    r_fonts.set(qn('w:hAnsi'), 'Calibri')
    r_fonts.set(qn('w:cs'), 'Arial')
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_title(text, size=22):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size, True, (20, 20, 20))
    return p


def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, 11, False, (90, 90, 90))
    return p


def add_h(text, level=1):
    p = doc.add_paragraph()
    sizes = {1: 15, 2: 12, 3: 11}
    color = (227, 24, 55) if level == 1 else (30, 30, 30)
    run = p.add_run(text)
    set_run_font(run, sizes.get(level, 12), True, color)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_p(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 11, bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    return p


def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(item)
        set_run_font(run, 11)
        p.paragraph_format.space_after = Pt(3)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, 10, True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ''
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, 10)
    doc.add_paragraph()


add_title('مستندات پروژه')
add_title('nevox news! — پلتفرم خبری دو زبانه', 18)
add_subtitle('پروژه پایانی | پلتفرم خبری Django — nevox news!')
add_subtitle('دانشجو: Setare Mousavi (ستاره موسوی)')
add_subtitle('GitHub: https://github.com/setarehmousavifar/news_platform.git')
add_p(
    'این فایل مخصوص ارسال ایمیلی به استاد تهیه شده است (فرمت Word). '
    'شامل توضیحات نوشتاری پروژه است و فایل اجرایی/EXE ندارد.'
)

add_h('۱. معرفی پروژه')
add_p(
    'nevox news یک پلتفرم خبری تحت وب است که با فریم‌ورک Django پیاده‌سازی شده. '
    'هدف پروژه، ساخت یک سامانه خبری واقعی برای ارائه دانشگاهی است؛ شامل نمایش اخبار، '
    'تعامل کاربران، پنل تحریریه، پشتیبانی دو زبانه (انگلیسی/فارسی)، API و مستندات نصب.'
)
add_p(
    'کاربر عادی می‌تواند خبر بخواند، لایک کند، نظر بدهد و خبر ذخیره کند. '
    'ادمین می‌تواند خبر منتشر کند. سوپرادمین مدیریت کاربران و کل محتوا را بر عهده دارد.'
)

add_h('۲. اهداف پروژه')
add_bullets([
    'طراحی و پیاده‌سازی یک سایت خبری کامل با رابط کاربری مدرن',
    'پیاده‌سازی سیستم نقش‌ها و سطوح دسترسی (RBAC)',
    'پشتیبانی دو زبانه انگلیسی و فارسی (UI + محتوا)',
    'فراهم‌کردن داده نمونه استاندارد برای ارزیابی نهایی',
    'ارائه REST API همراه با احراز هویت JWT و مستندات Swagger',
    'رعایت نکات امنیت، SEO و بهینه‌سازی سرعت',
])

add_h('۳. فناوری‌ها و ابزارها')
add_table(
    ['بخش', 'فناوری'],
    [
        ['Backend', 'Python 3.10+ / Django 4.2'],
        ['Database', 'MySQL / MariaDB (XAMPP برای توسعه محلی)'],
        ['Frontend', 'Django Templates + Bootstrap 5 + CSS/JS سفارشی'],
        ['API', 'Django REST Framework + SimpleJWT + drf-spectacular'],
        ['i18n', 'Django Locale + ذخیره ترجمه اخبار در JSONField'],
        ['Media', 'Pillow + Video URL/File'],
        ['Test', 'pytest + pytest-django'],
        ['Optional', 'Redis / Meilisearch / Docker'],
    ],
)

add_h('۴. ساختار کلی پروژه')
add_bullets([
    'accounts/ : کاربران، احراز هویت، نقش‌ها، تنظیمات سایت، داشبورد ادمین',
    'news/ : اخبار، دسته‌بندی، تگ/کلمات کلیدی، ترجمه، seed، API',
    'interactions/ : لایک، کامنت، ذخیره خبر (Saved stories)',
    'templates/ : قالب‌های HTML سمت سرور',
    'static/ : CSS، JS و دارایی‌های برند',
    'locale/fa/ : ترجمه‌های رابط کاربری فارسی',
    'news_platform/ : تنظیمات، URLها، middleware',
])

add_h('۵. امکانات اصلی سیستم')
add_h('۵.۱ بخش عمومی سایت', 2)
add_bullets([
    'صفحه اصلی مجله‌ای: Featured/Carousel، Trending، Popular، Latest، World News، Top Stories',
    'لیست اخبار با جستجو، فیلتر دسته، مرتب‌سازی و صفحه‌بندی',
    'صفحه جزئیات خبر: متن کامل، تصویر، کلمات کلیدی، ویدیو، اخبار مرتبط',
    'صفحات About و Contact',
    'تم روز/شب و پشتیبانی RTL برای فارسی',
    'SEO: meta، Open Graph، sitemap.xml، robots.txt، JSON-LD',
])

add_h('۵.۲ کاربران و نقش‌ها', 2)
add_table(
    ['نقش', 'دسترسی‌ها'],
    [
        ['normal', 'ثبت‌نام/ورود، خواندن خبر، لایک، کامنت، ذخیره خبر، پروفایل'],
        ['admin', 'دسترسی کاربر عادی + ساخت/ویرایش/بایگانی خبر خودش + داشبورد'],
        ['super_admin', 'مدیریت همه اخبار و کاربران + تغییر نقش + Django Admin'],
    ],
)

add_h('۵.۳ پنل تحریریه', 2)
add_bullets([
    'ساخت و ویرایش خبر با عنوان، متن، وضعیت انتشار، دسته، تصویر، ویدیو',
    'ورود Keywords/Tags به صورت ویرگول‌جدا',
    'تصاویر درون‌متنی با نشانه‌گذاری [inline:1]',
    'پیش‌نمایش (Preview) قبل از انتشار',
    'حذف نرم (Soft Delete) به جای حذف فیزیکی',
])

add_h('۵.۴ دو زبانه بودن', 2)
add_p(
    'رابط کاربری با سیستم i18n جنگو ترجمه می‌شود (فایل‌های locale/fa). '
    'متن اخبار به صورت انگلیسی حرفه‌ای ذخیره شده و ترجمه فارسی در فیلد translations '
    'نگه‌داری می‌شود تا هنگام سوییچ زبان، سایت کند نشود و به سرویس ترجمه در لحظه وابسته نباشد.'
)

add_h('۵.۵ API', 2)
add_bullets([
    'مستندات Swagger: /api/docs/',
    'JWT: /api/v1/auth/jwt/',
    'News / Categories / Comments تحت /api/v1/',
    'Health check: /health/',
])

add_h('۶. داده نمونه (Seed Data) چگونه وارد شده است؟')
add_p(
    'برای ارزیابی نهایی، اخبار به صورت دستی یکی‌یکی وارد نشده‌اند. '
    'یک Management Command به نام seed_premium پیاده‌سازی شده که دیتابیس را '
    'با داده استاندارد و قابل تکرار پر می‌کند.'
)
add_bullets([
    'حدود ۸۴ خبر انگلیسی ژورنالیستی',
    '۲۱ دسته‌بندی موضوعی سایت',
    'نویسنده، تاریخ انتشار، تگ، تعداد بازدید',
    'تصویر برای اخبار و ویدیو برای چند خبر خاص',
    'پر شدن بخش‌های صفحه اصلی مثل Trending، Latest و World',
])
add_p('دستور اجرا:')
add_p('python manage.py seed_premium --purge --skip-translate', bold=True)
add_p(
    'توضیح: --purge داده قبلی اخبار را پاک می‌کند و داده جدید می‌سازد. '
    'ترجمه فارسی اخبار اختیاری است و با دستور translate_news قابل اجراست (نیاز به اینترنت).'
)

add_h('۷. راهنمای نصب و اجرا برای استاد')
add_p('پیش‌نیاز: Python 3.10 تا 3.13 و MySQL/MariaDB (در ویندوز معمولاً XAMPP).')
add_bullets([
    'ساخت دیتابیس: CREATE DATABASE news_platform CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;',
    'ایجاد محیط مجازی: python -m venv venv',
    'فعال‌سازی: .\\venv\\Scripts\\Activate.ps1',
    'نصب وابستگی‌ها: pip install -r requirements.txt',
    'کپی تنظیمات: copy .env.example .env و تنظیم DB_PASSWORD',
    'مایگریشن: python manage.py migrate',
    'داده دمو: python manage.py seed_premium --purge --skip-translate',
    'اجرا: python manage.py runserver localhost:8000',
])
add_p('آدرس‌های مهم:')
add_bullets([
    'سایت انگلیسی: http://localhost:8000/',
    'سایت فارسی: http://localhost:8000/fa/',
    'Django Admin: http://localhost:8000/admin/',
    'API Docs: http://localhost:8000/api/docs/',
])
add_p('ورود دمو پس از seed:')
add_table(
    ['نقش', 'Username', 'Password'],
    [
        ['Super Admin', 'admin', 'admin1234'],
        ['Journalist (نمونه)', 'sarah_chen', 'seedpass123'],
    ],
)
add_p('نکته مهم: برای جلوگیری از خطای CSRF، آدرس را با localhost باز کنید و با 127.0.0.1 مخلوط نکنید.')
add_p('اگر نصب mysqlclient روی ویندوز خطا داد، PyMySQL کافی است؛ پروژه به‌صورت خودکار از آن استفاده می‌کند.')

add_h('۸. نکات امنیتی پیاده‌سازی‌شده')
add_bullets([
    'محافظت CSRF روی فرم‌های POST',
    'هش شدن رمز عبور توسط سیستم احراز هویت Django',
    'کنترل دسترسی بر اساس نقش (decorator/permission)',
    'Rate limit روی ثبت‌نام و ورود',
    'Soft-delete برای اخبار',
    'عدم ارسال فایل .env در تحویل پروژه',
])

add_h('۹. بهینه‌سازی و کیفیت')
add_bullets([
    'کش لیست‌های پرتکرار صفحه اصلی',
    'عدم فراخوانی ترجمه زنده در مسیر رندر صفحه',
    'جستجو با اولویت Meilisearch (اختیاری) سپس MySQL FULLTEXT',
    'تست‌های pytest برای بخش‌های مهم',
    'چک‌لیست دفاع و راهنمای نصب داخل پروژه',
])

add_h('۱۰. نحوه ارائه و ارزیابی پیشنهادی')
add_bullets([
    'مشاهده صفحه اصلی و بخش‌های خبری',
    'سوییچ زبان به فارسی و بررسی RTL',
    'ورود با نقش‌های مختلف و مقایسه دسترسی‌ها',
    'ساخت یک خبر آزمایشی همراه با Keywords',
    'مشاهده API Docs و یک درخواست نمونه',
])

add_h('۱۱. محدودیت‌ها و نکات صادقانه')
add_bullets([
    'رمزهای seed فقط برای محیط دمو/آموزشی هستند',
    'Redis و Meilisearch اختیاری‌اند و برای اجرای پایه لازم نیستند',
    'ترجمه فارسی متن اخبار به اینترنت وابسته است (در صورت اجرای translate_news)',
    'پروژه SPA جدا (مثل React) نیست؛ فرانت بر پایه Templateهای Django است و API برای توسعه آینده آماده شده',
])

add_h('۱۲. جمع‌بندی')
add_p(
    'پروژه nevox news یک سامانه خبری کامل برای ارائه پایانی است که همزمان '
    'رابط کاربری، نقش‌های کاربری، داده دموی استاندارد، دو زبانه بودن، امنیت پایه، '
    'SEO و API را پوشش می‌دهد. راهنمای نصب تکمیلی در فایل‌های README.md و INSTALL.txt نیز موجود است.'
)

add_p('')
add_subtitle('پایان مستندات — آماده ارسال ایمیلی برای استاد')
add_subtitle('Student: Setare Mousavi | Project: nevox news!')

desktop = Path(r'C:\Users\setar\OneDrive\Desktop\مستندات_پروژه_nevox_news_ستاره_موسوی.docx')
project_copy = Path(r'C:\Users\setar\OneDrive\Desktop\news_platform\docs\PROJECT_DOCUMENTATION_FA.docx')
project_copy.parent.mkdir(parents=True, exist_ok=True)
doc.save(desktop)
doc.save(project_copy)
print(f'Saved: {desktop}')
print(f'Saved: {project_copy}')
print(f'Size KB: {round(desktop.stat().st_size / 1024, 1)}')
